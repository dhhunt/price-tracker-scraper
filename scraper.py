"""Price tracker for books.toscrape.com.

Walks every catalogue page, collects each book's title, price, stock status
and star rating, removes duplicates, and saves the result to CSV and Excel.
Each run is added to a price history so you can see what changed since the
last run. Optionally pushes the latest prices to a Google Sheet.

books.toscrape.com is a public practice site built for scraping exercises.
"""

import argparse
import csv
import datetime as dt
import os
import time
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup

START_URL = "https://books.toscrape.com/catalogue/page-1.html"
USER_AGENT = "price-tracker-scraper (portfolio demo; contact via GitHub)"
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
FIELDS = ["title", "price_gbp", "in_stock", "rating", "url"]


def fetch(url, session=None, retries=3):
    """Download one page, retrying a couple of times on network hiccups."""
    session = session or requests.Session()
    for attempt in range(1, retries + 1):
        try:
            resp = session.get(url, headers={"User-Agent": USER_AGENT}, timeout=20)
            resp.raise_for_status()
            resp.encoding = "utf-8"  # the site is UTF-8; without this "£" turns into "Â£"
            return resp.text
        except requests.RequestException:
            if attempt == retries:
                raise
            time.sleep(2 * attempt)


def parse_page(html, page_url):
    """Return (books on this page, URL of the next page or None)."""
    soup = BeautifulSoup(html, "html.parser")
    books = []
    for pod in soup.select("article.product_pod"):
        link = pod.select_one("h3 a")
        price = pod.select_one("p.price_color").get_text(strip=True)
        stock = pod.select_one("p.availability").get_text(" ", strip=True)
        rating_tag = pod.select_one("p.star-rating")
        rating_word = next((c for c in rating_tag["class"] if c in RATINGS), None) if rating_tag else None
        books.append({
            "title": link.get("title") or link.get_text(strip=True),
            "price_gbp": float(price.replace("£", "").replace("Â", "").strip()),
            "in_stock": "in stock" in stock.lower(),
            "rating": RATINGS.get(rating_word),
            "url": urljoin(page_url, link["href"]),
        })
    next_link = soup.select_one("li.next a")
    next_url = urljoin(page_url, next_link["href"]) if next_link else None
    return books, next_url


def scrape(start_url=START_URL, max_pages=None, delay=1.0, get=fetch, log=print):
    """Follow the "next" button page by page and return de-duplicated books."""
    seen, books = set(), []
    url, page = start_url, 0
    while url and (max_pages is None or page < max_pages):
        page += 1
        html = get(url)
        page_books, url = parse_page(html, url)
        new = [b for b in page_books if b["url"] not in seen]
        seen.update(b["url"] for b in new)
        books.extend(new)
        log(f"page {page}: {len(page_books)} books ({len(new)} new), total {len(books)}")
        if url and delay:
            time.sleep(delay)  # be polite to the server
    return books


def load_history(path):
    if not os.path.exists(path):
        return []
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def update_history(history, books, run_date):
    """Add today's prices to the history (one row per book per day)."""
    kept = [r for r in history if r["date"] != run_date]  # re-running on the same day replaces that day
    for b in books:
        kept.append({"date": run_date, "title": b["title"], "url": b["url"], "price_gbp": f'{b["price_gbp"]:.2f}'})
    return kept


def price_changes(history, run_date):
    """Compare today's prices with each book's most recent earlier price."""
    previous = {}
    for r in sorted(history, key=lambda r: r["date"]):
        if r["date"] < run_date:
            previous[r["url"]] = r
    changes = []
    for r in history:
        if r["date"] != run_date or r["url"] not in previous:
            continue
        old, new = float(previous[r["url"]]["price_gbp"]), float(r["price_gbp"])
        if old != new:
            changes.append({
                "title": r["title"], "old_price": old, "new_price": new,
                "change": round(new - old, 2), "since": previous[r["url"]]["date"], "url": r["url"],
            })
    return changes


def write_csv(path, rows, fields):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def write_excel(path, books, changes):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    sheets = [("Latest prices", books, FIELDS),
              ("Price changes", changes, ["title", "old_price", "new_price", "change", "since", "url"])]
    for i, (name, rows, fields) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        ws.append(fields)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append([r[k] for k in fields])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            width = max(len(str(c.value)) for c in col if c.value is not None)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
    wb.save(path)


def push_to_google_sheet(sheet_name, creds_file, books):
    """Overwrite the first tab of a Google Sheet with the latest prices."""
    import gspread  # only needed for this optional step

    client = gspread.service_account(filename=creds_file)
    ws = client.open(sheet_name).sheet1
    ws.clear()
    ws.update([FIELDS] + [[b[k] for k in FIELDS] for b in books])


def main():
    p = argparse.ArgumentParser(description="Track book prices on books.toscrape.com")
    p.add_argument("--out-dir", default="output", help="where to save files (default: output)")
    p.add_argument("--max-pages", type=int, help="stop after this many pages (default: all 50)")
    p.add_argument("--delay", type=float, default=1.0, help="seconds to wait between pages (default: 1)")
    p.add_argument("--sheet", help="name of a Google Sheet to update (optional)")
    p.add_argument("--creds", default="service_account.json", help="Google service account key file")
    args = p.parse_args()

    os.makedirs(args.out_dir, exist_ok=True)
    run_date = dt.date.today().isoformat()

    books = scrape(max_pages=args.max_pages, delay=args.delay)

    history_path = os.path.join(args.out_dir, "price_history.csv")
    history = update_history(load_history(history_path), books, run_date)
    changes = price_changes(history, run_date)

    write_csv(os.path.join(args.out_dir, "latest_prices.csv"), books, FIELDS)
    write_csv(history_path, history, ["date", "title", "url", "price_gbp"])
    write_excel(os.path.join(args.out_dir, "prices.xlsx"), books, changes)
    print(f"Saved {len(books)} books to {args.out_dir}/ ({len(changes)} price changes since last run)")

    if args.sheet:
        push_to_google_sheet(args.sheet, args.creds, books)
        print(f'Updated Google Sheet "{args.sheet}"')


if __name__ == "__main__":
    main()
