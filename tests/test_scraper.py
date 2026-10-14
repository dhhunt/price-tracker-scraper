"""Offline tests: they read saved copies of two catalogue pages, so no internet is needed.

Run with:  python -m unittest discover tests
"""

import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import scraper  # noqa: E402

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures")
BASE = "https://books.toscrape.com/catalogue/"


def fake_get(url):
    name = url.rsplit("/", 1)[-1]
    with open(os.path.join(FIXTURES, name), encoding="utf-8") as f:
        return f.read()


class ScraperTests(unittest.TestCase):
    def scrape(self, **kw):
        return scraper.scrape(start_url=BASE + "page-1.html", delay=0, get=fake_get, log=lambda *_: None, **kw)

    def test_reads_every_field(self):
        books, next_url = scraper.parse_page(fake_get(BASE + "page-1.html"), BASE + "page-1.html")
        self.assertEqual(len(books), 3)
        self.assertEqual(books[0], {
            "title": "A Light in the Attic",
            "price_gbp": 51.77,
            "in_stock": True,
            "rating": 3,
            "url": BASE + "a-light-in-the-attic_1000/index.html",
        })
        self.assertEqual(next_url, BASE + "page-2.html")

    def test_follows_pagination_and_removes_duplicates(self):
        books = self.scrape()
        titles = [b["title"] for b in books]
        # "Soumission" appears on both pages but is kept once
        self.assertEqual(titles, ["A Light in the Attic", "Tipping the Velvet", "Soumission", "Sharp Objects"])

    def test_out_of_stock(self):
        # The saved page 2 was edited to mark "Sharp Objects" out of stock so this case is covered
        sharp = self.scrape()[-1]
        self.assertFalse(sharp["in_stock"])
        self.assertEqual(sharp["rating"], 4)

    def test_max_pages(self):
        self.assertEqual(len(self.scrape(max_pages=1)), 3)

    def test_price_changes_between_runs(self):
        books = self.scrape()
        history = scraper.update_history([], books, "2026-10-01")
        books[0] = dict(books[0], price_gbp=45.00)  # price drop on the second run
        history = scraper.update_history(history, books, "2026-10-08")
        changes = scraper.price_changes(history, "2026-10-08")
        self.assertEqual(len(changes), 1)
        self.assertEqual(changes[0]["title"], "A Light in the Attic")
        self.assertEqual(changes[0]["change"], -6.77)

    def test_same_day_rerun_replaces_that_day(self):
        books = self.scrape()
        history = scraper.update_history([], books, "2026-10-01")
        history = scraper.update_history(history, books, "2026-10-01")
        self.assertEqual(len(history), len(books))

    def test_writes_files(self):
        books = self.scrape()
        with tempfile.TemporaryDirectory() as d:
            scraper.write_csv(os.path.join(d, "latest.csv"), books, scraper.FIELDS)
            scraper.write_excel(os.path.join(d, "prices.xlsx"), books, [])
            from openpyxl import load_workbook
            wb = load_workbook(os.path.join(d, "prices.xlsx"))
            self.assertEqual(wb.sheetnames, ["Latest prices", "Price changes"])
            self.assertEqual(wb["Latest prices"].max_row, len(books) + 1)


if __name__ == "__main__":
    unittest.main()
